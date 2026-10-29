
from flask import (
    Flask, request, redirect, url_for, session, flash,
    render_template_string, g, send_file
)
import sqlite3
from functools import wraps
from datetime import date
from io import BytesIO
from werkzeug.security import generate_password_hash, check_password_hash
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

app = Flask(__name__)
app.secret_key = "replace-with-your-own-long-secret-key"
DB_NAME = "hospital.db"


# ==================== DATABASE ====================

def get_db():
    if "db" not in g:
        g.db = sqlite3.connect(DB_NAME)
        g.db.row_factory = sqlite3.Row
        g.db.execute("PRAGMA foreign_keys=ON")
    return g.db


@app.teardown_appcontext
def close_db(error=None):
    db = g.pop("db", None)
    if db:
        db.close()


def init_db():
    with app.app_context():
        db = get_db()
        db.executescript("""
        CREATE TABLE IF NOT EXISTS admins (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            username TEXT UNIQUE NOT NULL,
            password_hash TEXT NOT NULL
        );

        CREATE TABLE IF NOT EXISTS patients (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            age INTEGER NOT NULL,
            gender TEXT NOT NULL,
            phone TEXT NOT NULL,
            address TEXT,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        );

        CREATE TABLE IF NOT EXISTS doctors (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT NOT NULL,
            specialization TEXT NOT NULL,
            phone TEXT NOT NULL,
            fee REAL NOT NULL DEFAULT 0
        );

        CREATE TABLE IF NOT EXISTS appointments (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            patient_id INTEGER NOT NULL,
            doctor_id INTEGER NOT NULL,
            appointment_date TEXT NOT NULL,
            appointment_time TEXT NOT NULL,
            reason TEXT,
            status TEXT NOT NULL DEFAULT 'Scheduled',
            FOREIGN KEY(patient_id) REFERENCES patients(id) ON DELETE CASCADE,
            FOREIGN KEY(doctor_id) REFERENCES doctors(id) ON DELETE CASCADE
        );

        CREATE TABLE IF NOT EXISTS bills (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            patient_id INTEGER NOT NULL,
            description TEXT NOT NULL,
            amount REAL NOT NULL,
            bill_date TEXT NOT NULL,
            payment_status TEXT NOT NULL DEFAULT 'Unpaid',
            FOREIGN KEY(patient_id) REFERENCES patients(id) ON DELETE CASCADE
        );
        """)

        admin = db.execute(
            "SELECT id FROM admins WHERE username='admin'"
        ).fetchone()

        if not admin:
            db.execute(
                "INSERT INTO admins(username,password_hash) VALUES(?,?)",
                ("admin", generate_password_hash("Admin@123"))
            )

        db.commit()


# ==================== AUTHENTICATION ====================

def login_required(func):
    @wraps(func)
    def wrapper(*args, **kwargs):
        if "admin_id" not in session:
            return redirect(url_for("login"))
        return func(*args, **kwargs)
    return wrapper


# ==================== MODERN UI ====================

CSS = """
<style>
@import url('https://fonts.googleapis.com/css2?family=Inter:wght@400;500;600;700;800&display=swap');

:root {
    --navy: #14263d;
    --navy2: #1d3854;
    --teal: #0d9488;
    --teal-dark: #0f766e;
    --bg: #f4f7fb;
    --text: #26364a;
    --muted: #77869a;
    --border: #e7edf4;
    --white: #fff;
}

* { box-sizing: border-box; }

body {
    margin: 0;
    background: var(--bg);
    color: var(--text);
    font-family: Inter, "Segoe UI", Arial, sans-serif;
    font-size: 14px;
    line-height: 1.6;
}

a { text-decoration: none; }

.app-shell {
    display: flex;
    min-height: 100vh;
}

.sidebar {
    width: 245px;
    flex-shrink: 0;
    background: linear-gradient(180deg, #14263d, #102a43);
    color: white;
    padding: 23px 15px;
    display: flex;
    flex-direction: column;
}

.brand {
    display: flex;
    align-items: center;
    gap: 11px;
    padding: 5px 8px 25px;
    border-bottom: 1px solid #ffffff20;
    margin-bottom: 23px;
}

.brand-icon {
    display: grid;
    place-items: center;
    width: 43px;
    height: 43px;
    background: linear-gradient(135deg, #2dd4bf, #0d9488);
    border-radius: 13px;
    font-size: 23px;
}

.brand h2 {
    color: white;
    font-size: 16px;
    line-height: 1.4;
    margin: 0;
}

.brand small {
    color: #a8bdcf;
    font-size: 10px;
    letter-spacing: 1px;
}

.menu-label {
    padding: 0 12px;
    color: #91a6bb;
    font-size: 10px;
    font-weight: 800;
    letter-spacing: 1.5px;
}

.side-nav {
    display: flex;
    flex-direction: column;
    gap: 7px;
}

.side-nav a {
    color: #c7d4e2;
    display: flex;
    align-items: center;
    gap: 12px;
    padding: 12px;
    border-radius: 10px;
    font-weight: 600;
    transition: .2s;
}

.side-nav a:hover,
.side-nav a.active {
    background: linear-gradient(100deg, #0d9488, #0f766e);
    color: white;
    transform: translateX(2px);
}

.side-icon {
    width: 23px;
    text-align: center;
    font-size: 17px;
}

.sidebar-bottom { margin-top: auto; padding-top: 22px; }

.admin-mini {
    background: #ffffff10;
    border: 1px solid #ffffff20;
    padding: 12px;
    border-radius: 11px;
    margin-bottom: 12px;
    overflow-wrap: anywhere;
}

.admin-mini strong { display: block; color: white; }
.admin-mini small { color: #a8bdcf; }

.main-area { flex: 1; min-width: 0; }

.topbar {
    min-height: 80px;
    padding: 17px 32px;
    display: flex;
    align-items: center;
    justify-content: space-between;
    gap: 12px;
    background: white;
    border-bottom: 1px solid var(--border);
}

.topbar h2 { color: var(--navy); font-size: 19px; margin: 0; }
.topbar p { margin: 3px 0 0; color: var(--muted); font-size: 12px; }

.top-pill {
    background: #e3f8f4;
    color: #0f766e;
    border-radius: 25px;
    padding: 8px 12px;
    font-size: 11px;
    font-weight: 700;
    white-space: nowrap;
}

main {
    max-width: 1600px;
    margin: auto;
    padding: 29px 32px 45px;
}

h1 {
    color: var(--navy);
    font-size: clamp(24px, 3vw, 31px);
    letter-spacing: -.8px;
    font-weight: 800;
    margin: 0 0 23px;
}

h2, h3 { color: var(--navy); }

h3 {
    font-size: 16px;
    font-weight: 750;
    margin: 0 0 18px;
}

.card {
    background: white;
    border: 1px solid var(--border);
    border-radius: 16px;
    padding: 23px;
    margin-bottom: 22px;
    box-shadow: 0 7px 25px #14263d09;
    overflow-x: auto;
}

.card h3 {
    border-bottom: 1px solid #edf1f6;
    padding-bottom: 13px;
}

.stats {
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(175px, 1fr));
    gap: 16px;
    margin-bottom: 25px;
}

.stat {
    background: white;
    border: 1px solid var(--border);
    border-left: 4px solid var(--teal);
    border-radius: 15px;
    padding: 21px;
    box-shadow: 0 7px 23px #14263d08;
    transition: .2s;
}

.stat:hover {
    transform: translateY(-4px);
    box-shadow: 0 13px 28px #14263d12;
}

.stat p {
    color: var(--muted);
    margin: 0 0 6px;
    font-weight: 600;
}

.stat h2 {
    margin: 0;
    font-size: 34px;
    color: var(--teal-dark);
    letter-spacing: -1px;
}

.grid {
    display: grid;
    grid-template-columns: repeat(auto-fit, minmax(190px, 1fr));
    gap: 13px;
}

label {
    display: block;
    color: #46566b;
    font-weight: 650;
    margin-top: 7px;
}

input, select {
    display: block;
    width: 100%;
    max-width: 100%;
    min-height: 43px;
    padding: 10px 12px;
    margin: 7px 0 13px;
    border: 1px solid #dce5ef;
    border-radius: 9px;
    background: #fbfdff;
    color: var(--text);
    font: inherit;
    outline: none;
}

input:focus, select:focus {
    border-color: var(--teal);
    background: white;
    box-shadow: 0 0 0 3px #0d94881a;
}

button, .btn {
    display: inline-block;
    background: linear-gradient(135deg, #0d9488, #0f766e);
    color: white;
    border: 0;
    border-radius: 9px;
    padding: 10px 14px;
    margin: 3px 3px 3px 0;
    font-family: inherit;
    font-size: 12px;
    font-weight: 750;
    cursor: pointer;
    transition: .2s;
}

button:hover, .btn:hover {
    transform: translateY(-2px);
    box-shadow: 0 6px 13px #0d94882b;
}

button.delete {
    background: #fff0f0;
    color: #c62828;
    border: 1px solid #ffd8d8;
    box-shadow: none;
}

button.delete:hover { background: #fee2e2; }

table {
    width: 100%;
    min-width: 650px;
    border-collapse: separate;
    border-spacing: 0;
    font-size: 12px;
}

th {
    background: #f0f7f8;
    color: #45616b;
    font-size: 10px;
    letter-spacing: .6px;
    text-transform: uppercase;
    font-weight: 800;
}

th, td {
    padding: 13px 11px;
    text-align: left;
    border-bottom: 1px solid #edf1f6;
    vertical-align: middle;
}

td { color: #4a5b70; }
tr:last-child td { border-bottom: 0; }
table tr:hover td { background: #f8fcfc; }
form.inline { display: inline; }

.msg {
    background: #e6f8f0;
    color: #126747;
    border: 1px solid #c5eddb;
    border-radius: 10px;
    padding: 12px 15px;
    margin-bottom: 17px;
    font-weight: 600;
}

.small { color: var(--muted); font-size: 12px; }

/* Login screen */
.login-screen {
    min-height: 100vh;
    display: grid;
    place-items: center;
    padding: 22px;
    background:
        radial-gradient(circle at 15% 15%, #ccfbf1, transparent 32%),
        radial-gradient(circle at 90% 90%, #dbeafe, transparent 35%),
        #f3f7fb;
}

.login-card {
    width: 100%;
    max-width: 420px;
    background: white;
    padding: 35px;
    border-radius: 22px;
    border: 1px solid var(--border);
    box-shadow: 0 24px 65px #14263d14;
}

.login-logo {
    width: 62px;
    height: 62px;
    display: grid;
    place-items: center;
    background: linear-gradient(135deg, #0d9488, #14263d);
    border-radius: 17px;
    font-size: 30px;
    margin-bottom: 20px;
}

.login-card h1 { font-size: 25px; margin-bottom: 4px; }
.login-card .subtitle { color: var(--muted); margin: 0 0 23px; }
.login-card button { width: 100%; min-height: 45px; margin-top: 8px; }
.login-note {
    margin-top: 18px;
    padding: 10px;
    border-radius: 8px;
    background: #f0fdfa;
    color: #0f766e;
    font-size: 11px;
    text-align: center;
}

@media(max-width: 900px) {
    .sidebar { width: 215px; }
    main { padding: 23px 18px; }
    .topbar { padding: 16px 18px; }
}

@media(max-width: 650px) {
    .app-shell { display: block; }
    .sidebar { width: 100%; padding: 13px; }
    .brand { padding-bottom: 13px; margin-bottom: 10px; }
    .menu-label { display: none; }
    .side-nav { display: grid; grid-template-columns: repeat(2, 1fr); gap: 5px; }
    .side-nav a { padding: 9px; font-size: 12px; }
    .sidebar-bottom { padding-top: 10px; }
    .admin-mini { margin-bottom: 8px; }
    .topbar { min-height: 68px; }
    main { padding: 20px 12px; }
    .stats { grid-template-columns: repeat(2, minmax(0, 1fr)); gap: 10px; }
    .stat { padding: 14px; }
    .stat h2 { font-size: 26px; }
    .card { padding: 15px; }
    .grid { grid-template-columns: 1fr; }
}
</style>
"""


def page(title, content, public=False, **data):
    if public:
        template = """
        <!doctype html><html><head><meta charset="utf-8">
        <meta name="viewport" content="width=device-width,initial-scale=1">
        <title>{{ title }} | MediCare HMS</title>
        """ + CSS + """</head><body>
        {% for m in get_flashed_messages() %}
        <div class="msg" style="position:fixed;top:12px;left:12px;right:12px;z-index:5">{{ m }}</div>
        {% endfor %}
        """ + content + """
        </body></html>
        """
        return render_template_string(template, title=title, **data)

    items = [
        ("/", "📊", "Dashboard"),
        ("/patients", "🧑‍🤝‍🧑", "Patients"),
        ("/doctors", "🩺", "Doctors"),
        ("/appointments", "📅", "Appointments"),
        ("/billing", "💳", "Billing"),
        ("/export-excel", "📈", "Excel Reports"),
    ]

    nav = ""
    for href, emoji, label in items:
        active = current_active(request.path, href)
        cls = "active" if active else ""
        nav += (
            f'<a class="{cls}" href="{href}">'
            f'<span class="side-icon">{emoji}</span>{label}</a>'
        )

    template = """
    <!doctype html><html><head><meta charset="utf-8">
    <meta name="viewport" content="width=device-width,initial-scale=1">
    <title>{{ title }} | MediCare HMS</title>
    """ + CSS + """</head><body>
    <div class="app-shell">
        <aside class="sidebar">
            <div class="brand">
                <div class="brand-icon">🏥</div>
                <div><h2>MediCare HMS</h2><small>HOSPITAL MANAGEMENT</small></div>
            </div>
            <p class="menu-label">MAIN MENU</p>
            <nav class="side-nav">""" + nav + """</nav>
            <div class="sidebar-bottom">
                <div class="admin-mini">
                    <strong>👤 {{ session.get('username', 'Admin') }}</strong>
                    <small>Administrator</small>
                </div>
                <form method="post" action="/logout">
                    <button class="logout-btn">↪ &nbsp; Logout</button>
                </form>
            </div>
        </aside>
        <section class="main-area">
            <header class="topbar">
                <div><h2>{{ title }}</h2>
                <p>Welcome to your hospital management workspace.</p></div>
                <div class="top-pill">● System Dashboard</div>
            </header>
            <main>
                {% for m in get_flashed_messages() %}<div class="msg">{{ m }}</div>{% endfor %}
                """ + content + """
            </main>
        </section>
    </div></body></html>
    """
    return render_template_string(template, title=title, **data)


def current_active(path, href):
    return path == href or (href != "/" and path.startswith(href))


# ==================== LOGIN / LOGOUT ====================

LOGIN_HTML = """
<div class="login-screen">
    <div class="login-card">
        <div class="login-logo">🏥</div>
        <h1>MediCare Hospital</h1>
        <p class="subtitle">Sign in to your admin dashboard</p>
        <form method="post">
            <label>Username</label>
            <input name="username" placeholder="Enter username" required>
            <label>Password</label>
            <input name="password" type="password" placeholder="Enter password" required>
            <button type="submit">Sign In &nbsp; →</button>
        </form>
        <div class="login-note">Hospital Management System</div>
    </div>
</div>
"""


@app.route("/login", methods=["GET", "POST"])
def login():
    if request.method == "POST":
        user = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        admin = get_db().execute(
            "SELECT * FROM admins WHERE username=?", (user,)
        ).fetchone()

        if admin and check_password_hash(admin["password_hash"], password):
            session.clear()
            session["admin_id"] = admin["id"]
            session["username"] = admin["username"]
            return redirect(url_for("dashboard"))

        flash("Invalid login details.")

    return page("Admin Login", LOGIN_HTML, public=True)


@app.route("/logout", methods=["POST"])
@login_required
def logout():
    session.clear()
    return redirect(url_for("login"))


# ==================== DASHBOARD ====================

@app.route("/")
@login_required
def dashboard():
    db = get_db()
    counts = {
        "Patients": db.execute("SELECT COUNT(*) FROM patients").fetchone()[0],
        "Doctors": db.execute("SELECT COUNT(*) FROM doctors").fetchone()[0],
        "Appointments": db.execute("SELECT COUNT(*) FROM appointments").fetchone()[0],
        "Bills": db.execute("SELECT COUNT(*) FROM bills").fetchone()[0],
    }
    recent = db.execute("""
        SELECT a.*, p.name AS patient, d.name AS doctor
        FROM appointments a
        JOIN patients p ON p.id=a.patient_id
        JOIN doctors d ON d.id=a.doctor_id
        ORDER BY a.appointment_date DESC, a.appointment_time DESC
        LIMIT 10
    """).fetchall()

    content = """
    <h1>Good to see you, Admin 👋</h1>
    <div class="stats">
    {% for k,v in counts.items() %}
        <div class="stat"><p>{{ k }}</p><h2>{{ v }}</h2></div>
    {% endfor %}
    </div>
    <div class="card"><h3>📅 Recent Appointments</h3>
    <table><tr><th>Patient</th><th>Doctor</th><th>Date</th><th>Time</th><th>Status</th></tr>
    {% for a in recent %}
    <tr><td>{{ a.patient }}</td><td>{{ a.doctor }}</td>
    <td>{{ a.appointment_date }}</td><td>{{ a.appointment_time }}</td><td>{{ a.status }}</td></tr>
    {% else %}<tr><td colspan="5">No appointments yet.</td></tr>{% endfor %}
    </table></div>
    """
    return page("Dashboard", content, counts=counts, recent=recent)


# ==================== PATIENTS ====================

@app.route("/patients")
@login_required
def patients():
    q = request.args.get("q", "").strip()
    records = get_db().execute("""
        SELECT * FROM patients
        WHERE name LIKE ? OR phone LIKE ?
        ORDER BY id DESC
    """, (f"%{q}%", f"%{q}%")).fetchall()

    content = """
    <h1>Patient Management</h1>
    <div class="card"><h3>➕ Register New Patient</h3>
    <form method="post" action="/patients/add"><div class="grid">
    <input name="name" placeholder="Full name" required>
    <input name="age" type="number" min="0" max="125" placeholder="Age" required>
    <select name="gender" required><option value="">Select gender</option>
    <option>Male</option><option>Female</option><option>Other</option></select>
    <input name="phone" placeholder="Phone number" required>
    <input name="address" placeholder="Address">
    </div><button>Save Patient</button></form></div>

    <div class="card"><h3>👥 Patient Records</h3>
    <form method="get"><input name="q" value="{{ q }}" placeholder="Search by name or phone">
    <button>Search</button></form>
    <table><tr><th>ID</th><th>Name</th><th>Age</th><th>Gender</th>
    <th>Phone</th><th>Address</th><th>Actions</th></tr>
    {% for p in records %}<tr>
    <td>{{ p.id }}</td><td>{{ p.name }}</td><td>{{ p.age }}</td>
    <td>{{ p.gender }}</td><td>{{ p.phone }}</td><td>{{ p.address or '' }}</td>
    <td><a class="btn" href="/patients/edit/{{ p.id }}">Edit</a>
    <a class="btn" href="/patients/pdf/{{ p.id }}">📄 Record PDF</a>
    <form class="inline" method="post" action="/patients/delete/{{ p.id }}"
    onsubmit="return confirm('Delete patient?')">
    <button class="delete">Delete</button></form></td></tr>
    {% else %}<tr><td colspan="7">No patients found.</td></tr>{% endfor %}
    </table></div>
    """
    return page("Patients", content, records=records, q=q)


@app.route("/patients/add", methods=["POST"])
@login_required
def add_patient():
    try:
        name = request.form.get("name", "").strip()
        age = int(request.form.get("age", ""))
        gender = request.form.get("gender", "")
        phone = request.form.get("phone", "").strip()
        address = request.form.get("address", "").strip()

        if not name or not phone or gender not in ["Male", "Female", "Other"]:
            raise ValueError
        if not 0 <= age <= 125:
            raise ValueError

        db = get_db()
        db.execute("""
            INSERT INTO patients(name,age,gender,phone,address)
            VALUES(?,?,?,?,?)
        """, (name, age, gender, phone, address))
        db.commit()
        flash("Patient added successfully.")
    except ValueError:
        flash("Please enter valid patient details.")
    return redirect(url_for("patients"))


@app.route("/patients/edit/<int:pid>", methods=["GET", "POST"])
@login_required
def edit_patient(pid):
    db = get_db()
    p = db.execute("SELECT * FROM patients WHERE id=?", (pid,)).fetchone()
    if not p:
        flash("Patient not found.")
        return redirect(url_for("patients"))

    if request.method == "POST":
        try:
            name = request.form.get("name", "").strip()
            age = int(request.form.get("age", ""))
            gender = request.form.get("gender", "")
            phone = request.form.get("phone", "").strip()
            address = request.form.get("address", "").strip()

            if not name or not phone or gender not in ["Male", "Female", "Other"]:
                raise ValueError
            if not 0 <= age <= 125:
                raise ValueError

            db.execute("""
                UPDATE patients SET name=?,age=?,gender=?,phone=?,address=?
                WHERE id=?
            """, (name, age, gender, phone, address, pid))
            db.commit()
            flash("Patient updated.")
            return redirect(url_for("patients"))
        except ValueError:
            flash("Please enter valid details.")

    content = """
    <div class="card"><h3>Edit Patient</h3>
    <form method="post"><div class="grid">
    <input name="name" value="{{ p.name }}" required>
    <input name="age" type="number" min="0" max="125" value="{{ p.age }}" required>
    <select name="gender">{% for g in ['Male','Female','Other'] %}
    <option {{ 'selected' if p.gender==g else '' }}>{{ g }}</option>
    {% endfor %}</select>
    <input name="phone" value="{{ p.phone }}" required>
    <input name="address" value="{{ p.address or '' }}">
    </div><button>Update Patient</button></form></div>
    """
    return page("Edit Patient", content, p=p)


@app.route("/patients/delete/<int:pid>", methods=["POST"])
@login_required
def delete_patient(pid):
    db = get_db()
    db.execute("DELETE FROM patients WHERE id=?", (pid,))
    db.commit()
    flash("Patient deleted.")
    return redirect(url_for("patients"))


# ==================== PATIENT RECORD PDF ====================

@app.route("/patients/pdf/<int:pid>")
@login_required
def download_patient_pdf(pid):
    db = get_db()
    patient = db.execute(
        "SELECT * FROM patients WHERE id=?", (pid,)
    ).fetchone()

    if not patient:
        flash("Patient not found.")
        return redirect(url_for("patients"))

    appointments = db.execute("""
        SELECT a.appointment_date, a.appointment_time, a.reason, a.status,
               d.name AS doctor
        FROM appointments a
        JOIN doctors d ON d.id=a.doctor_id
        WHERE a.patient_id=?
        ORDER BY a.appointment_date DESC, a.appointment_time DESC
    """, (pid,)).fetchall()

    bills = db.execute("""
        SELECT description, amount, bill_date, payment_status
        FROM bills WHERE patient_id=? ORDER BY id DESC
    """, (pid,)).fetchall()

    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    _, height = A4
    y = height - 50

    def write_line(text, size=10, bold=False):
        nonlocal y
        if y < 60:
            pdf.showPage()
            y = height - 50
        pdf.setFont("Helvetica-Bold" if bold else "Helvetica", size)
        pdf.drawString(50, y, str(text)[:100])
        y -= size + 8

    pdf.setTitle(f"Patient Record {patient['id']}")
    write_line("MEDICARE HOSPITAL", 18, True)
    write_line("INDIVIDUAL PATIENT RECORD", 14, True)
    write_line("-" * 55)
    write_line(f"Patient ID: {patient['id']}")
    write_line(f"Name: {patient['name']}")
    write_line(f"Age: {patient['age']}")
    write_line(f"Gender: {patient['gender']}")
    write_line(f"Phone: {patient['phone']}")
    write_line(f"Address: {patient['address'] or 'Not provided'}")
    write_line(f"Registered: {patient['created_at'] or 'N/A'}")

    y -= 10
    write_line("APPOINTMENT HISTORY", 13, True)
    if appointments:
        for a in appointments:
            write_line(f"Date: {a['appointment_date']}  Time: {a['appointment_time']}")
            write_line(f"Doctor: {a['doctor']}")
            write_line(f"Reason: {a['reason'] or 'N/A'}")
            write_line(f"Status: {a['status']}")
            y -= 4
    else:
        write_line("No appointments found.")

    y -= 10
    write_line("BILLING HISTORY", 13, True)
    total = 0.0
    if bills:
        for b in bills:
            amount = float(b["amount"])
            total += amount
            write_line(f"Date: {b['bill_date']}")
            write_line(f"Description: {b['description']}")
            write_line(f"Amount: INR {amount:.2f}")
            write_line(f"Payment Status: {b['payment_status']}")
            y -= 4
    else:
        write_line("No bills found.")

    y -= 10
    write_line(f"Total billed: INR {total:.2f}", 12, True)
    pdf.save()
    output.seek(0)

    return send_file(
        output,
        as_attachment=True,
        download_name=f"Patient_Record_{patient['id']}.pdf",
        mimetype="application/pdf"
    )


# ==================== DOCTORS ====================

@app.route("/doctors")
@login_required
def doctors():
    q = request.args.get("q", "").strip()
    records = get_db().execute("""
        SELECT * FROM doctors
        WHERE name LIKE ? OR specialization LIKE ? OR phone LIKE ?
        ORDER BY id DESC
    """, (f"%{q}%", f"%{q}%", f"%{q}%")).fetchall()

    content = """
    <h1>Doctor Management</h1>
    <div class="card"><h3>➕ Add Doctor</h3>
    <form method="post" action="/doctors/add"><div class="grid">
    <input name="name" placeholder="Doctor name" required>
    <input name="specialization" placeholder="Specialization" required>
    <input name="phone" placeholder="Phone" required>
    <input name="fee" type="number" min="0" step="0.01" placeholder="Consultation fee" required>
    </div><button>Add Doctor</button></form></div>

    <div class="card"><h3>🩺 Doctor Directory</h3>
    <form method="get"><input name="q" value="{{ q }}" placeholder="Search doctor, specialty or phone">
    <button>Search</button></form>
    <table><tr><th>ID</th><th>Name</th><th>Specialization</th><th>Phone</th><th>Fee</th><th>Actions</th></tr>
    {% for d in records %}<tr><td>{{ d.id }}</td><td>{{ d.name }}</td>
    <td>{{ d.specialization }}</td><td>{{ d.phone }}</td><td>₹{{ '%.2f'|format(d.fee) }}</td>
    <td><a class="btn" href="/doctors/edit/{{ d.id }}">Edit</a>
    <form class="inline" method="post" action="/doctors/delete/{{ d.id }}"
    onsubmit="return confirm('Delete doctor?')"><button class="delete">Delete</button></form>
    </td></tr>{% else %}<tr><td colspan="6">No doctors found.</td></tr>{% endfor %}
    </table></div>
    """
    return page("Doctors", content, records=records, q=q)


@app.route("/doctors/add", methods=["POST"])
@login_required
def add_doctor():
    try:
        name = request.form.get("name", "").strip()
        specialization = request.form.get("specialization", "").strip()
        phone = request.form.get("phone", "").strip()
        fee = float(request.form.get("fee", ""))

        if not name or not specialization or not phone or fee < 0:
            raise ValueError

        db = get_db()
        db.execute("""
            INSERT INTO doctors(name,specialization,phone,fee)
            VALUES(?,?,?,?)
        """, (name, specialization, phone, fee))
        db.commit()
        flash("Doctor added successfully.")
    except ValueError:
        flash("Enter valid doctor details.")
    return redirect(url_for("doctors"))


@app.route("/doctors/edit/<int:did>", methods=["GET", "POST"])
@login_required
def edit_doctor(did):
    db = get_db()
    d = db.execute("SELECT * FROM doctors WHERE id=?", (did,)).fetchone()
    if not d:
        flash("Doctor not found.")
        return redirect(url_for("doctors"))

    if request.method == "POST":
        try:
            name = request.form.get("name", "").strip()
            specialization = request.form.get("specialization", "").strip()
            phone = request.form.get("phone", "").strip()
            fee = float(request.form.get("fee", ""))

            if not name or not specialization or not phone or fee < 0:
                raise ValueError

            db.execute("""
                UPDATE doctors SET name=?,specialization=?,phone=?,fee=?
                WHERE id=?
            """, (name, specialization, phone, fee, did))
            db.commit()
            flash("Doctor updated.")
            return redirect(url_for("doctors"))
        except ValueError:
            flash("Enter valid doctor details.")

    content = """
    <div class="card"><h3>Edit Doctor</h3><form method="post">
    <input name="name" value="{{ d.name }}" required>
    <input name="specialization" value="{{ d.specialization }}" required>
    <input name="phone" value="{{ d.phone }}" required>
    <input name="fee" type="number" min="0" step="0.01" value="{{ d.fee }}" required>
    <button>Update Doctor</button></form></div>
    """
    return page("Edit Doctor", content, d=d)


@app.route("/doctors/delete/<int:did>", methods=["POST"])
@login_required
def delete_doctor(did):
    db = get_db()
    db.execute("DELETE FROM doctors WHERE id=?", (did,))
    db.commit()
    flash("Doctor deleted.")
    return redirect(url_for("doctors"))


# ==================== APPOINTMENTS ====================

@app.route("/appointments")
@login_required
def appointments():
    db = get_db()
    records = db.execute("""
        SELECT a.*, p.name AS patient, d.name AS doctor
        FROM appointments a
        JOIN patients p ON p.id=a.patient_id
        JOIN doctors d ON d.id=a.doctor_id
        ORDER BY a.appointment_date DESC, a.appointment_time DESC
    """).fetchall()
    patients_list = db.execute("SELECT * FROM patients ORDER BY name").fetchall()
    doctors_list = db.execute("SELECT * FROM doctors ORDER BY name").fetchall()

    content = """
    <h1>Appointments</h1>
    <div class="card"><h3>📅 Book Appointment</h3>
    {% if patients and doctors %}
    <form method="post" action="/appointments/add"><div class="grid">
    <select name="patient_id" required><option value="">Select patient</option>
    {% for p in patients %}<option value="{{ p.id }}">{{ p.name }}</option>{% endfor %}</select>
    <select name="doctor_id" required><option value="">Select doctor</option>
    {% for d in doctors %}<option value="{{ d.id }}">{{ d.name }} - {{ d.specialization }}</option>{% endfor %}</select>
    <input name="appointment_date" type="date" min="{{ today }}" required>
    <input name="appointment_time" type="time" required>
    <input name="reason" placeholder="Reason for visit">
    </div><button>Book Appointment</button></form>
    {% else %}<p>Add a patient and doctor first.</p>{% endif %}</div>

    <div class="card"><h3>Appointment List</h3><table>
    <tr><th>ID</th><th>Patient</th><th>Doctor</th><th>Date</th><th>Time</th><th>Reason</th><th>Status</th><th>Actions</th></tr>
    {% for a in records %}<tr><td>{{ a.id }}</td><td>{{ a.patient }}</td><td>{{ a.doctor }}</td>
    <td>{{ a.appointment_date }}</td><td>{{ a.appointment_time }}</td><td>{{ a.reason or '' }}</td><td>{{ a.status }}</td>
    <td><form method="post" action="/appointments/update/{{ a.id }}">
    <select name="status">{% for s in ['Scheduled','Completed','Cancelled'] %}
    <option {{ 'selected' if a.status==s else '' }}>{{ s }}</option>{% endfor %}</select>
    <button>Save</button></form>
    <form method="post" action="/appointments/delete/{{ a.id }}" onsubmit="return confirm('Delete appointment?')">
    <button class="delete">Delete</button></form></td></tr>
    {% else %}<tr><td colspan="8">No appointments.</td></tr>{% endfor %}
    </table></div>
    """
    return page(
        "Appointments", content, records=records,
        patients=patients_list, doctors=doctors_list,
        today=date.today().isoformat()
    )


@app.route("/appointments/add", methods=["POST"])
@login_required
def add_appointment():
    db = get_db()
    try:
        pid = int(request.form.get("patient_id", ""))
        did = int(request.form.get("doctor_id", ""))
        appt_date = request.form.get("appointment_date", "")
        appt_time = request.form.get("appointment_time", "")
        reason = request.form.get("reason", "").strip()
        date.fromisoformat(appt_date)

        if appt_date < date.today().isoformat() or not appt_time:
            raise ValueError
        if not db.execute("SELECT id FROM patients WHERE id=?", (pid,)).fetchone():
            raise ValueError
        if not db.execute("SELECT id FROM doctors WHERE id=?", (did,)).fetchone():
            raise ValueError

        db.execute("""
            INSERT INTO appointments(patient_id,doctor_id,appointment_date,
            appointment_time,reason) VALUES(?,?,?,?,?)
        """, (pid, did, appt_date, appt_time, reason))
        db.commit()
        flash("Appointment booked.")
    except (ValueError, TypeError, sqlite3.IntegrityError):
        db.rollback()
        flash("Invalid appointment details.")
    return redirect(url_for("appointments"))


@app.route("/appointments/update/<int:aid>", methods=["POST"])
@login_required
def update_appointment(aid):
    status = request.form.get("status", "")
    if status not in ["Scheduled", "Completed", "Cancelled"]:
        flash("Invalid status.")
    else:
        db = get_db()
        db.execute("UPDATE appointments SET status=? WHERE id=?", (status, aid))
        db.commit()
        flash("Appointment status updated.")
    return redirect(url_for("appointments"))


@app.route("/appointments/delete/<int:aid>", methods=["POST"])
@login_required
def delete_appointment(aid):
    db = get_db()
    db.execute("DELETE FROM appointments WHERE id=?", (aid,))
    db.commit()
    flash("Appointment deleted.")
    return redirect(url_for("appointments"))


# ==================== BILLING ====================

@app.route("/billing", methods=["GET", "POST"])
@login_required
def billing():
    db = get_db()

    if request.method == "POST":
        try:
            pid = int(request.form.get("patient_id", ""))
            description = request.form.get("description", "").strip()
            amount = float(request.form.get("amount", ""))
            bill_date = request.form.get("bill_date", "") or date.today().isoformat()
            date.fromisoformat(bill_date)

            patient = db.execute(
                "SELECT id FROM patients WHERE id=?", (pid,)
            ).fetchone()

            if not patient or not description or amount <= 0:
                raise ValueError

            db.execute("""
                INSERT INTO bills(patient_id,description,amount,bill_date)
                VALUES(?,?,?,?)
            """, (pid, description, amount, bill_date))
            db.commit()
            flash("Bill created successfully.")
        except (ValueError, TypeError):
            db.rollback()
            flash("Enter valid bill details.")
        return redirect(url_for("billing"))

    patients_list = db.execute("SELECT * FROM patients ORDER BY name").fetchall()
    bills = db.execute("""
        SELECT b.*, p.name AS patient
        FROM bills b JOIN patients p ON p.id=b.patient_id
        ORDER BY b.id DESC
    """).fetchall()

    content = """
    <h1>Billing System</h1>
    <div class="card"><h3>💳 Create Bill</h3>
    {% if patients %}
    <form method="post"><div class="grid">
    <select name="patient_id" required><option value="">Select patient</option>
    {% for p in patients %}<option value="{{ p.id }}">{{ p.name }}</option>{% endfor %}</select>
    <input name="description" placeholder="Consultation / Service" required>
    <input name="amount" type="number" min="0.01" step="0.01" placeholder="Amount ₹" required>
    <input name="bill_date" type="date" value="{{ today }}" required>
    </div><button>Create Bill</button></form>
    {% else %}<p>Add a patient before creating a bill.</p>{% endif %}</div>

    <div class="card"><h3>Bill Records</h3><table>
    <tr><th>Bill ID</th><th>Patient</th><th>Description</th><th>Amount</th><th>Date</th><th>Payment</th><th>Actions</th></tr>
    {% for b in bills %}<tr><td>{{ b.id }}</td><td>{{ b.patient }}</td>
    <td>{{ b.description }}</td><td>₹{{ '%.2f'|format(b.amount) }}</td>
    <td>{{ b.bill_date }}</td><td>{{ b.payment_status }}</td><td>
    <form method="post" action="/billing/pay/{{ b.id }}">
    <button>{{ 'Mark Unpaid' if b.payment_status=='Paid' else 'Mark Paid' }}</button></form>
    <a class="btn" href="/billing/pdf/{{ b.id }}">📄 PDF</a>
    </td></tr>{% else %}<tr><td colspan="7">No bills yet.</td></tr>{% endfor %}
    </table></div>
    """
    return page(
        "Billing", content, patients=patients_list, bills=bills,
        today=date.today().isoformat()
    )


@app.route("/billing/pay/<int:bid>", methods=["POST"])
@login_required
def toggle_bill(bid):
    db = get_db()
    bill = db.execute("SELECT * FROM bills WHERE id=?", (bid,)).fetchone()
    if bill:
        status = "Unpaid" if bill["payment_status"] == "Paid" else "Paid"
        db.execute("UPDATE bills SET payment_status=? WHERE id=?", (status, bid))
        db.commit()
        flash("Payment status updated.")
    return redirect(url_for("billing"))


# ==================== BILL PDF ====================

@app.route("/billing/pdf/<int:bid>")
@login_required
def download_bill_pdf(bid):
    db = get_db()
    bill = db.execute("""
        SELECT b.*, p.name AS patient, p.phone, p.address
        FROM bills b JOIN patients p ON p.id=b.patient_id
        WHERE b.id=?
    """, (bid,)).fetchone()

    if not bill:
        flash("Bill not found.")
        return redirect(url_for("billing"))

    output = BytesIO()
    pdf = canvas.Canvas(output, pagesize=A4)
    width, height = A4

    pdf.setTitle(f"Hospital Bill {bill['id']}")
    pdf.setFont("Helvetica-Bold", 20)
    pdf.drawString(50, height-60, "MediCare Hospital")
    pdf.setFont("Helvetica", 11)
    pdf.drawString(50, height-85, "Hospital Billing Invoice")
    pdf.line(50, height-100, width-50, height-100)

    y = height-140
    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, f"Bill Number: {bill['id']}")
    y -= 25
    pdf.setFont("Helvetica", 11)
    pdf.drawString(50, y, f"Bill Date: {bill['bill_date']}")
    y -= 40

    pdf.setFont("Helvetica-Bold", 12)
    pdf.drawString(50, y, "Patient Details")
    y -= 25
    pdf.setFont("Helvetica", 11)
    pdf.drawString(50, y, f"Name: {bill['patient']}")
    y -= 20
    pdf.drawString(50, y, f"Phone: {bill['phone']}")
    y -= 20
    pdf.drawString(50, y, f"Address: {(bill['address'] or 'Not provided')[:75]}")
    y -= 45

    pdf.line(50, y, width-50, y)
    y -= 28
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(50, y, "Description")
    pdf.drawString(350, y, "Amount")
    y -= 25
    pdf.setFont("Helvetica", 11)
    pdf.drawString(50, y, bill["description"][:45])
    pdf.drawRightString(width-50, y, f"INR {bill['amount']:.2f}")
    y -= 30
    pdf.line(50, y, width-50, y)
    y -= 28
    pdf.setFont("Helvetica-Bold", 13)
    pdf.drawRightString(width-50, y, f"Total: INR {bill['amount']:.2f}")
    y -= 28
    pdf.setFont("Helvetica", 11)
    pdf.drawString(50, y, f"Payment Status: {bill['payment_status']}")
    y -= 55
    pdf.drawString(50, y, "Thank you for choosing MediCare Hospital.")

    pdf.save()
    output.seek(0)
    return send_file(
        output, as_attachment=True,
        download_name=f"Hospital_Bill_{bill['id']}.pdf",
        mimetype="application/pdf"
    )


# ==================== EXCEL REPORTS ====================

@app.route("/export-excel")
@login_required
def export_excel():
    db = get_db()
    wb = Workbook()

    sheets = [
        ("Patients", "SELECT id,name,age,gender,phone,address FROM patients"),
        ("Doctors", "SELECT id,name,specialization,phone,fee FROM doctors"),
        ("Appointments", """
            SELECT a.id,p.name,d.name,a.appointment_date,a.appointment_time,
                   a.reason,a.status
            FROM appointments a
            JOIN patients p ON p.id=a.patient_id
            JOIN doctors d ON d.id=a.doctor_id
        """),
        ("Bills", """
            SELECT b.id,p.name,b.description,b.amount,b.bill_date,b.payment_status
            FROM bills b JOIN patients p ON p.id=b.patient_id
        """)
    ]

    for index, (title, query) in enumerate(sheets):
        ws = wb.active if index == 0 else wb.create_sheet()
        ws.title = title
        cursor = db.execute(query)
        ws.append([col[0] for col in cursor.description])

        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="0F766E")

        for row in cursor.fetchall():
            ws.append(list(row))

        ws.freeze_panes = "A2"
        for column in ws.columns:
            letter = column[0].column_letter
            max_len = max(len(str(c.value or "")) for c in column)
            ws.column_dimensions[letter].width = min(max_len+3, 35)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return send_file(
        output, as_attachment=True,
        download_name="Hospital_Report.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


# ==================== START APP ====================

if __name__ == "__main__":
    init_db()
    app.run(debug=True)
