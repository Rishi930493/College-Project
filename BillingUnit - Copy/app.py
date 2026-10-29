from flask import Flask, render_template, request, redirect, jsonify, send_file
import sqlite3
import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import (
    SimpleDocTemplate,
    Table,
    TableStyle,
    Paragraph,
    Spacer
)
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm


app = Flask(__name__)


# =========================
# DATABASE CONNECTION
# =========================

def get_database():
    connection = sqlite3.connect("database.db")
    connection.row_factory = sqlite3.Row
    return connection


# =========================
# DASHBOARD
# =========================

@app.route("/")
def home():

    connection = get_database()

    today = datetime.now().strftime("%Y-%m-%d")
    current_month = datetime.now().strftime("%Y-%m")

    # Today's Bills
    today_bills = connection.execute("""
        SELECT COUNT(*) AS count
        FROM bills
        WHERE substr(bill_date, 1, 10) = ?
    """, (today,)).fetchone()["count"]

    # Today's Collection
    today_sales = connection.execute("""
        SELECT COALESCE(SUM(total), 0) AS total
        FROM bills
        WHERE substr(bill_date, 1, 10) = ?
    """, (today,)).fetchone()["total"]

    # Monthly Collection
    monthly_sales = connection.execute("""
        SELECT COALESCE(SUM(total), 0) AS total
        FROM bills
        WHERE substr(bill_date, 1, 7) = ?
    """, (current_month,)).fetchone()["total"]

    # Customer Count
    customer_count = connection.execute("""
        SELECT COUNT(DISTINCT customer_mobile) AS count
        FROM bills
        WHERE customer_mobile IS NOT NULL
        AND TRIM(customer_mobile) != ''
    """).fetchone()["count"]

    connection.close()

    return render_template(
        "dashboard.html",
        today_bills=today_bills,
        today_sales=today_sales,
        monthly_sales=monthly_sales,
        customer_count=customer_count
    )


# =========================
# PRODUCTS
# =========================

@app.route("/products")
def products():

    connection = get_database()

    products = connection.execute(
        "SELECT * FROM products ORDER BY id DESC"
    ).fetchall()

    connection.close()

    return render_template(
        "products.html",
        products=products
    )


@app.route("/add-product", methods=["POST"])
def add_product():

    name = request.form["name"]
    price = request.form["price"]

    connection = get_database()

    connection.execute(
        "INSERT INTO products (name, price) VALUES (?, ?)",
        (name, price)
    )

    connection.commit()
    connection.close()

    return redirect("/products")


@app.route("/delete-product/<int:product_id>")
def delete_product(product_id):

    connection = get_database()

    connection.execute(
        "DELETE FROM products WHERE id = ?",
        (product_id,)
    )

    connection.commit()
    connection.close()

    return redirect("/products")


# =========================
# BILLING PAGE
# =========================

@app.route("/billing")
def billing():

    connection = get_database()

    products = connection.execute(
        "SELECT * FROM products ORDER BY name"
    ).fetchall()

    result = connection.execute("""
        SELECT COALESCE(MAX(bill_no), 0) + 1
        AS next_bill_no
        FROM bills
    """).fetchone()

    next_bill_no = result["next_bill_no"]

    connection.close()

    return render_template(
        "billing.html",
        products=products,
        next_bill_no=next_bill_no
    )


# =========================
# SAVE BILL
# =========================

@app.route("/save-bill", methods=["POST"])
def save_bill():

    data = request.get_json()

    customer_name = data.get("customer_name", "")
    customer_mobile = data.get("customer_mobile", "")
    payment_method = data.get("payment_method", "Cash")
    total = data.get("total", 0)
    items = data.get("items", [])

    if not items:
        return jsonify({
            "success": False,
            "message": "Please add at least one item."
        })

    connection = get_database()

    try:

        result = connection.execute("""
            SELECT COALESCE(MAX(bill_no), 0) + 1
            AS next_bill_no
            FROM bills
        """).fetchone()

        bill_no = result["next_bill_no"]

        bill_date = datetime.now().strftime(
            "%Y-%m-%d %H:%M:%S"
        )

        cursor = connection.execute("""
            INSERT INTO bills
            (
                bill_no,
                bill_date,
                total,
                payment_method,
                customer_name,
                customer_mobile
            )
            VALUES (?, ?, ?, ?, ?, ?)
        """, (
            bill_no,
            bill_date,
            total,
            payment_method,
            customer_name,
            customer_mobile
        ))

        bill_id = cursor.lastrowid

        for item in items:

            connection.execute("""
                INSERT INTO bill_items
                (
                    bill_id,
                    product_name,
                    quantity,
                    price,
                    total
                )
                VALUES (?, ?, ?, ?, ?)
            """, (
                bill_id,
                item["name"],
                item["quantity"],
                item["price"],
                item["total"]
            ))

        connection.commit()

        return jsonify({
            "success": True,
            "bill_no": bill_no,
            "message": "Bill saved successfully!"
        })

    except Exception as error:

        connection.rollback()

        return jsonify({
            "success": False,
            "message": str(error)
        })

    finally:

        connection.close()


# =========================
# DAILY ENTRIES
# =========================

@app.route("/daily-entries")
def daily_entries():

    connection = get_database()

    today = datetime.now().strftime("%Y-%m-%d")

    bills = connection.execute("""
        SELECT *
        FROM bills
        ORDER BY id DESC
    """).fetchall()

    today_bills = connection.execute("""
        SELECT COUNT(*) AS count
        FROM bills
        WHERE substr(bill_date, 1, 10) = ?
    """, (today,)).fetchone()["count"]

    today_sales = connection.execute("""
        SELECT COALESCE(SUM(total), 0) AS total
        FROM bills
        WHERE substr(bill_date, 1, 10) = ?
    """, (today,)).fetchone()["total"]

    connection.close()

    return render_template(
        "daily_entries.html",
        bills=bills,
        today_bills=today_bills,
        today_sales=today_sales
    )


# =========================
# VIEW BILL
# =========================

@app.route("/view-bill/<int:bill_id>")
def view_bill(bill_id):

    connection = get_database()

    bill = connection.execute("""
        SELECT *
        FROM bills
        WHERE id = ?
    """, (bill_id,)).fetchone()

    if bill is None:

        connection.close()

        return "Bill not found"

    items = connection.execute("""
        SELECT *
        FROM bill_items
        WHERE bill_id = ?
        ORDER BY id
    """, (bill_id,)).fetchall()

    # Shop Settings
    shop_settings = connection.execute("""
        SELECT *
        FROM settings
        WHERE id = 1
    """).fetchone()

    connection.close()

    return render_template(
        "view_bill.html",
        bill=bill,
        items=items,
        settings=shop_settings
    )


# =========================
# REPORTS
# =========================

@app.route("/reports")
def reports():

    connection = get_database()

    today = datetime.now().strftime("%Y-%m-%d")

    # Date filter
    from_date = request.args.get(
        "from_date",
        today
    )

    to_date = request.args.get(
        "to_date",
        today
    )

    # Date range condition
    start_datetime = from_date + " 00:00:00"
    end_datetime = to_date + " 23:59:59"

    # Total Sales
    total_sales = connection.execute("""
        SELECT COALESCE(SUM(total), 0) AS total
        FROM bills
        WHERE bill_date BETWEEN ? AND ?
    """, (
        start_datetime,
        end_datetime
    )).fetchone()["total"]

    # Total Bills
    total_bills = connection.execute("""
        SELECT COUNT(*) AS count
        FROM bills
        WHERE bill_date BETWEEN ? AND ?
    """, (
        start_datetime,
        end_datetime
    )).fetchone()["count"]

    # Cash Collection
    cash_sales = connection.execute("""
        SELECT COALESCE(SUM(total), 0) AS total
        FROM bills
        WHERE bill_date BETWEEN ?
        AND ?
        AND payment_method = 'Cash'
    """, (
        start_datetime,
        end_datetime
    )).fetchone()["total"]

    # UPI Collection
    upi_sales = connection.execute("""
        SELECT COALESCE(SUM(total), 0) AS total
        FROM bills
        WHERE bill_date BETWEEN ?
        AND ?
        AND payment_method = 'UPI'
    """, (
        start_datetime,
        end_datetime
    )).fetchone()["total"]

    # Bills for selected date range
    bills = connection.execute("""
        SELECT *
        FROM bills
        WHERE bill_date BETWEEN ? AND ?
        ORDER BY id DESC
    """, (
        start_datetime,
        end_datetime
    )).fetchall()

    connection.close()

    return render_template(
        "reports.html",
        total_sales=total_sales,
        total_bills=total_bills,
        cash_sales=cash_sales,
        upi_sales=upi_sales,
        bills=bills,
        from_date=from_date,
        to_date=to_date
    )


# =========================
# EXPORT EXCEL
# =========================

@app.route("/export-excel")
def export_excel():

    connection = get_database()

    today = datetime.now().strftime("%Y-%m-%d")

    from_date = request.args.get(
        "from_date",
        today
    )

    to_date = request.args.get(
        "to_date",
        today
    )

    start_datetime = from_date + " 00:00:00"
    end_datetime = to_date + " 23:59:59"

    bills = connection.execute("""
        SELECT *
        FROM bills
        WHERE bill_date BETWEEN ? AND ?
        ORDER BY id
    """, (
        start_datetime,
        end_datetime
    )).fetchall()

    connection.close()

    # Summary

    total_sales = sum(
        float(bill["total"] or 0)
        for bill in bills
    )

    cash_sales = sum(
        float(bill["total"] or 0)
        for bill in bills
        if bill["payment_method"] == "Cash"
    )

    upi_sales = sum(
        float(bill["total"] or 0)
        for bill in bills
        if bill["payment_method"] == "UPI"
    )

    total_bills = len(bills)

    # Create workbook

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "Sales Report"

    # Title

    sheet["A1"] = "Food Shop Sales Report"

    sheet["A1"].font = Font(
        bold=True,
        size=16
    )

    sheet["A2"] = (
        f"From: {from_date}    "
        f"To: {to_date}"
    )

    sheet["A2"].font = Font(
        bold=True
    )

    # Summary

    sheet["A4"] = "Summary"

    sheet["A4"].font = Font(
        bold=True,
        size=13
    )

    sheet["A5"] = "Total Bills"
    sheet["B5"] = total_bills

    sheet["A6"] = "Total Sales"
    sheet["B6"] = total_sales

    sheet["A7"] = "Cash Sales"
    sheet["B7"] = cash_sales

    sheet["A8"] = "UPI Sales"
    sheet["B8"] = upi_sales

    # Table headings

    headers = [
        "Bill No",
        "Date & Time",
        "Customer",
        "Mobile",
        "Total",
        "Payment Method"
    ]

    header_row = 10

    for column, header in enumerate(
        headers,
        start=1
    ):

        cell = sheet.cell(
            row=header_row,
            column=column,
            value=header
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # Bill data

    row_number = header_row + 1

    for bill in bills:

        sheet.cell(
            row=row_number,
            column=1,
            value=bill["bill_no"]
        )

        sheet.cell(
            row=row_number,
            column=2,
            value=bill["bill_date"]
        )

        sheet.cell(
            row=row_number,
            column=3,
            value=bill["customer_name"] or ""
        )

        sheet.cell(
            row=row_number,
            column=4,
            value=bill["customer_mobile"] or ""
        )

        sheet.cell(
            row=row_number,
            column=5,
            value=float(
                bill["total"] or 0
            )
        )

        sheet.cell(
            row=row_number,
            column=6,
            value=bill["payment_method"]
        )

        row_number += 1

    # Total row

    total_row = row_number + 1

    sheet.cell(
        row=total_row,
        column=4,
        value="TOTAL"
    ).font = Font(
        bold=True
    )

    sheet.cell(
        row=total_row,
        column=5,
        value=total_sales
    ).font = Font(
        bold=True
    )

    # Currency format

    for row in range(6, 9):

        sheet.cell(
            row=row,
            column=2
        ).number_format = '₹#,##0.00'

    for row in range(
        header_row + 1,
        total_row + 1
    ):

        sheet.cell(
            row=row,
            column=5
        ).number_format = '₹#,##0.00'

    # Column widths

    sheet.column_dimensions["A"].width = 12
    sheet.column_dimensions["B"].width = 22
    sheet.column_dimensions["C"].width = 25
    sheet.column_dimensions["D"].width = 18
    sheet.column_dimensions["E"].width = 15
    sheet.column_dimensions["F"].width = 18

    # Save in memory

    output = io.BytesIO()

    workbook.save(output)

    output.seek(0)

    filename = (
        f"sales_report_"
        f"{from_date}_to_{to_date}.xlsx"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )


# =========================
# EXPORT PDF
# =========================

@app.route("/export-pdf")
def export_pdf():

    connection = get_database()

    today = datetime.now().strftime("%Y-%m-%d")

    from_date = request.args.get(
        "from_date",
        today
    )

    to_date = request.args.get(
        "to_date",
        today
    )

    start_datetime = from_date + " 00:00:00"
    end_datetime = to_date + " 23:59:59"

    bills = connection.execute("""
        SELECT *
        FROM bills
        WHERE bill_date BETWEEN ? AND ?
        ORDER BY id
    """, (
        start_datetime,
        end_datetime
    )).fetchall()

    connection.close()

    # Summary

    total_bills = len(bills)

    total_sales = sum(
        float(bill["total"] or 0)
        for bill in bills
    )

    cash_sales = sum(
        float(bill["total"] or 0)
        for bill in bills
        if bill["payment_method"] == "Cash"
    )

    upi_sales = sum(
        float(bill["total"] or 0)
        for bill in bills
        if bill["payment_method"] == "UPI"
    )

    # PDF memory file

    output = io.BytesIO()

    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        rightMargin=12 * mm,
        leftMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=12 * mm
    )

    styles = getSampleStyleSheet()

    elements = []

    # Title

    title = Paragraph(
        "<b>Food Shop Sales Report</b>",
        styles["Title"]
    )

    elements.append(title)

    # Date range

    date_text = Paragraph(
        f"<b>From:</b> {from_date} "
        f"&nbsp;&nbsp;&nbsp; "
        f"<b>To:</b> {to_date}",
        styles["Normal"]
    )

    elements.append(date_text)

    elements.append(
        Spacer(1, 8)
    )

    # Summary table

    summary_data = [
        [
            "Total Bills",
            "Total Sales",
            "Cash Sales",
            "UPI Sales"
        ],
        [
            str(total_bills),
            f"Rs. {total_sales:.2f}",
            f"Rs. {cash_sales:.2f}",
            f"Rs. {upi_sales:.2f}"
        ]
    ]

    summary_table = Table(
        summary_data,
        colWidths=[
            55 * mm,
            55 * mm,
            55 * mm,
            55 * mm
        ]
    )

    summary_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.lightgrey
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.black
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "FONTNAME",
                (0, 1),
                (-1, 1),
                "Helvetica-Bold"
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, -1),
                "CENTER"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                7
            )
        ])
    )

    elements.append(summary_table)

    elements.append(
        Spacer(1, 12)
    )

    # Bill table

    table_data = [
        [
            "Bill No",
            "Date & Time",
            "Customer",
            "Mobile",
            "Total",
            "Payment"
        ]
    ]

    for bill in bills:

        table_data.append([
            str(bill["bill_no"]),
            str(bill["bill_date"]),
            bill["customer_name"] or "",
            bill["customer_mobile"] or "",
            f"Rs. {float(bill['total'] or 0):.2f}",
            bill["payment_method"]
        ])

    # No bills

    if not bills:

        table_data.append([
            "",
            "",
            "No bills found",
            "",
            "",
            ""
        ])

    bill_table = Table(
        table_data,
        colWidths=[
            25 * mm,
            45 * mm,
            55 * mm,
            40 * mm,
            35 * mm,
            35 * mm
        ],
        repeatRows=1
    )

    bill_table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.darkgrey
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold"
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, 0),
                "CENTER"
            ),
            (
                "ALIGN",
                (0, 1),
                (0, -1),
                "CENTER"
            ),
            (
                "ALIGN",
                (4, 1),
                (5, -1),
                "CENTER"
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.5,
                colors.grey
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE"
            ),
            (
                "PADDING",
                (0, 0),
                (-1, -1),
                6
            )
        ])
    )

    elements.append(bill_table)

    elements.append(
        Spacer(1, 12)
    )

    # Final total

    total_text = Paragraph(
        f"<b>Total Sales: Rs. {total_sales:.2f}</b>",
        styles["Heading3"]
    )

    elements.append(total_text)

    # Create PDF

    document.build(elements)

    output.seek(0)

    filename = (
        f"sales_report_"
        f"{from_date}_to_{to_date}.pdf"
    )

    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/pdf"
    )


# =========================
# CUSTOMERS
# =========================

@app.route("/customers")
def customers():

    connection = get_database()

    customers = connection.execute("""
        SELECT
            customer_name,
            customer_mobile,
            COUNT(*) AS total_bills,
            COALESCE(SUM(total), 0) AS total_purchase
        FROM bills
        WHERE customer_mobile IS NOT NULL
        AND TRIM(customer_mobile) != ''
        GROUP BY customer_mobile
        ORDER BY MAX(id) DESC
    """).fetchall()

    connection.close()

    return render_template(
        "customers.html",
        customers=customers
    )


# =========================
# SETTINGS
# =========================

@app.route("/settings")
def settings():

    connection = get_database()

    settings_data = connection.execute("""
        SELECT *
        FROM settings
        WHERE id = 1
    """).fetchone()

    connection.close()

    return render_template(
        "settings.html",
        settings=settings_data
    )


@app.route("/save-settings", methods=["POST"])
def save_settings():

    shop_name = request.form.get(
        "shop_name",
        ""
    ).strip()

    mobile = request.form.get(
        "mobile",
        ""
    ).strip()

    address = request.form.get(
        "address",
        ""
    ).strip()

    bill_header = request.form.get(
        "bill_header",
        ""
    ).strip()

    connection = get_database()

    existing = connection.execute("""
        SELECT id
        FROM settings
        WHERE id = 1
    """).fetchone()

    if existing:

        connection.execute("""
            UPDATE settings
            SET shop_name = ?,
                mobile = ?,
                address = ?,
                bill_header = ?
            WHERE id = 1
        """, (
            shop_name,
            mobile,
            address,
            bill_header
        ))

    else:

        connection.execute("""
            INSERT INTO settings
            (
                id,
                shop_name,
                mobile,
                address,
                bill_header
            )
            VALUES (1, ?, ?, ?, ?)
        """, (
            shop_name,
            mobile,
            address,
            bill_header
        ))

    connection.commit()

    connection.close()

    return redirect("/settings")


# =========================
# RUN APP
# =========================

if __name__ == "__main__":
    app.run(debug=True)